"""NICTA Telecommunications QoS Rule 2022: Appendix A/B tables (spec v3 M16).

Inputs are OpenNetIQ exports with campaign tags (region, area_class,
location_name): `*_voice_sms.csv` (Schedule 1), `*_qoe.csv` (FTP and web)
and `*_tests.csv` (ping). Every rate is computed from raw rows, rounded to
one decimal place as the Rule's templates require.
Rule traceability: docs/standards/NICTA-QOS-RULE-2022.md.
"""

from __future__ import annotations

import csv
import json
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

from .io import Row, parse_float, parse_time, read_csv

RULE = "nicta-qos-2022"
UNTAGGED = "(untagged)"

# Schedule 1
CALL_FAILURE_MAX_PCT = 2.0
CALL_DROP_MAX_PCT = 2.0
MIN_CALLS_PER_REGION = 100
CALL_FAILED = {"not_placed", "rejected", "no_answer", "setup_timeout"}
CALL_ESTABLISHED = {"completed", "dropped"}
CALL_UNKNOWN = {"unknown", "ended_early"}

# Schedule 2
FTP_DL_MIN_MBPS = 2.5
FTP_UL_MIN_MBPS = 0.5
PING_MAX_MS = 250.0
PING_SHARE_PCT = 80.0
WEB_MAX_S = 10.0
WEB_SHARE_PCT = 80.0
LEVEL_SHARE_PCT = 80.0
RSRP_MIN_DBM = -110.0
RSCP_MIN_DBM = -100.0
MIN_FTP_PER_LOCATION = 10
MIN_WEB_PER_LOCATION = 10
MIN_PINGS_PER_LOCATION = 100
MIN_LOCATIONS_PER_REGION = 30


def ftp_share_pct(year: int | None) -> float:
    """80 % of samples in 2022-2023, 90 % from 2024 (Schedule 2)."""
    return 80.0 if year is not None and year < 2024 else 90.0


def pct(part: int, whole: int) -> float | None:
    return None if whole == 0 else round(100.0 * part / whole, 1)


def mean(values: list[float]) -> float | None:
    return None if not values else round(sum(values) / len(values), 1)


@dataclass
class Period:
    """A Rule measurement period: 1 Jan - 30 Jun (H1) or 1 Jul - 31 Dec (H2)."""

    label: str
    start: datetime
    end: datetime

    @staticmethod
    def parse(text: str) -> Period:
        t = text.strip().upper()
        if len(t) != 6 or t[4] != "H" or t[5] not in "12" or not t[:4].isdigit():
            raise ValueError(f"period must look like 2026H1 or 2026H2, got '{text}'")
        year = int(t[:4])
        if t[5] == "1":
            return Period(t, datetime(year, 1, 1, tzinfo=UTC), datetime(year, 7, 1, tzinfo=UTC))
        return Period(t, datetime(year, 7, 1, tzinfo=UTC), datetime(year + 1, 1, 1, tzinfo=UTC))

    def contains(self, t: datetime | None) -> bool:
        return t is not None and self.start <= t < self.end


@dataclass
class Exports:
    voice: list[Row] = field(default_factory=list)
    qoe: list[Row] = field(default_factory=list)
    tests: list[Row] = field(default_factory=list)
    files: list[str] = field(default_factory=list)


def find_exports(paths: Iterable[Path]) -> Exports:
    """Loads export CSVs from files or directories (recursive)."""
    out = Exports()
    seen: set[Path] = set()
    for p in paths:
        if not p.exists():
            raise ValueError(f"not found: {p}")
        candidates = sorted(p.rglob("*.csv")) if p.is_dir() else [p]
        for f in candidates:
            if f in seen:
                continue
            seen.add(f)
            name = f.name
            if name.endswith("_voice_sms.csv"):
                out.voice += read_csv(f)
            elif name.endswith("_qoe.csv"):
                out.qoe += read_csv(f)
            elif name.endswith("_tests.csv"):
                out.tests += read_csv(f)
            else:
                continue
            out.files.append(str(f))
    if not out.files:
        raise ValueError("no *_voice_sms.csv, *_qoe.csv or *_tests.csv exports found")
    return out


def _key(row: Row) -> tuple[str, str]:
    region = (row.get("region") or "").strip() or UNTAGGED
    area = (row.get("area_class") or "").strip() or UNTAGGED
    return region, area


def _location(row: Row) -> str:
    return (row.get("location_name") or "").strip() or UNTAGGED


def _in_period(rows: list[Row], period: Period | None) -> list[Row]:
    if period is None:
        return rows
    return [r for r in rows if period.contains(parse_time(r.get("timestamp")))]


@dataclass
class VoiceRow:
    """Appendix A for one region / area class."""

    region: str
    area_class: str
    attempts: int = 0
    failed: int = 0
    established: int = 0
    dropped: int = 0
    unknown: int = 0
    setup_times_s: list[float] = field(default_factory=list)
    technologies: dict[str, int] = field(default_factory=dict)

    @property
    def failure_rate_pct(self) -> float | None:
        return pct(self.failed, self.attempts)

    @property
    def drop_rate_pct(self) -> float | None:
        return pct(self.dropped, self.established)

    @property
    def failure_ok(self) -> bool | None:
        r = self.failure_rate_pct
        return None if r is None else r <= CALL_FAILURE_MAX_PCT

    @property
    def drop_ok(self) -> bool | None:
        r = self.drop_rate_pct
        return None if r is None else r <= CALL_DROP_MAX_PCT

    @property
    def mean_setup_s(self) -> float | None:
        return mean(self.setup_times_s)


def appendix_a(rows: list[Row]) -> list[VoiceRow]:
    groups: dict[tuple[str, str], VoiceRow] = {}
    for r in rows:
        if (r.get("kind") or "") != "call":
            continue
        k = _key(r)
        g = groups.setdefault(k, VoiceRow(*k))
        outcome = (r.get("outcome") or "").strip()
        if outcome in CALL_UNKNOWN:
            g.unknown += 1
            continue
        g.attempts += 1
        if outcome in CALL_FAILED:
            g.failed += 1
        if outcome in CALL_ESTABLISHED:
            g.established += 1
            setup = parse_float(r.get("setup_ms"))
            if setup is not None:
                g.setup_times_s.append(setup / 1000)
        if outcome == "dropped":
            g.dropped += 1
        tech = (r.get("voice_rat") or "").strip() or "unknown"
        g.technologies[tech] = g.technologies.get(tech, 0) + 1
    return [groups[k] for k in sorted(groups)]


@dataclass
class Share:
    """Samples meeting a threshold, with the mean of the measured values."""

    samples: int = 0
    meeting: int = 0
    values: list[float] = field(default_factory=list)
    required_pct: float = 80.0

    @property
    def share_pct(self) -> float | None:
        return pct(self.meeting, self.samples)

    @property
    def mean(self) -> float | None:
        return mean(self.values)

    @property
    def ok(self) -> bool | None:
        s = self.share_pct
        return None if s is None else s >= self.required_pct

    def add(self, value: float | None, meets: bool) -> None:
        self.samples += 1
        if value is not None:
            self.values.append(value)
        if meets:
            self.meeting += 1


@dataclass
class BroadbandRow:
    """Appendix B for one region / area class."""

    region: str
    area_class: str
    ftp_dl: Share = field(default_factory=Share)
    ftp_ul: Share = field(default_factory=Share)
    ping: Share = field(default_factory=Share)
    web: Share = field(default_factory=Share)
    rsrp: Share = field(default_factory=Share)
    rscp: Share = field(default_factory=Share)
    locations: set[str] = field(default_factory=set)


def appendix_b(qoe: list[Row], tests: list[Row], ftp_pct: float) -> list[BroadbandRow]:
    groups: dict[tuple[str, str], BroadbandRow] = {}

    def group(r: Row) -> BroadbandRow:
        k = _key(r)
        if k not in groups:
            g = BroadbandRow(*k)
            g.ftp_dl.required_pct = g.ftp_ul.required_pct = ftp_pct
            g.ping.required_pct = PING_SHARE_PCT
            g.web.required_pct = WEB_SHARE_PCT
            g.rsrp.required_pct = g.rscp.required_pct = LEVEL_SHARE_PCT
            groups[k] = g
        return groups[k]

    for r in qoe:
        kind = r.get("kind")
        ok = r.get("status") == "ok"
        if kind == "file" and r.get("protocol") == "ftp":
            g = group(r)
            g.locations.add(_location(r))
            rate = parse_float(r.get("rate_mbps"))
            if r.get("direction") == "upload":
                meets = ok and rate is not None and rate >= FTP_UL_MIN_MBPS
                g.ftp_ul.add(rate if ok else None, meets)
                continue
            meets = ok and rate is not None and rate >= FTP_DL_MIN_MBPS
            g.ftp_dl.add(rate if ok else None, meets)
            level = parse_float(r.get("serving_level_dbm"))
            rat = (r.get("serving_rat") or "").upper()
            if level is not None and rat == "LTE":
                g.rsrp.add(level, level >= RSRP_MIN_DBM)
            elif level is not None and rat == "WCDMA":
                g.rscp.add(level, level >= RSCP_MIN_DBM)
        elif kind == "web":
            g = group(r)
            g.locations.add(_location(r))
            total = parse_float(r.get("total_ms"))
            seconds = None if total is None or not ok else total / 1000
            g.web.add(seconds, seconds is not None and seconds <= WEB_MAX_S)
    for r in tests:
        if r.get("test_type") != "latency" or r.get("method") != "icmp":
            continue
        g = group(r)
        g.locations.add(_location(r))
        try:
            rtts = json.loads(r.get("raw_rtts_ms") or "[]")
        except json.JSONDecodeError:
            rtts = []
        for rtt in rtts if isinstance(rtts, list) else []:
            value = float(rtt) if isinstance(rtt, int | float) else None
            g.ping.add(value, value is not None and value <= PING_MAX_MS)
    return [groups[k] for k in sorted(groups)]


@dataclass
class SampleCheck:
    scope: str
    item: str
    required: int
    actual: int

    @property
    def ok(self) -> bool:
        return self.actual >= self.required


def sample_checks(voice: list[VoiceRow], qoe: list[Row], tests: list[Row]) -> list[SampleCheck]:
    checks: list[SampleCheck] = []
    calls: dict[str, int] = defaultdict(int)
    for v in voice:
        calls[v.region] += v.attempts
    for region in sorted(calls):
        item = "call samples (Sch. 1 G11)"
        checks.append(SampleCheck(region, item, MIN_CALLS_PER_REGION, calls[region]))
    per_loc: dict[tuple[str, str], dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for r in qoe:
        loc = (_key(r)[0], _location(r))
        if r.get("kind") == "file" and r.get("protocol") == "ftp":
            per_loc[loc]["ftp_upload" if r.get("direction") == "upload" else "ftp_download"] += 1
        elif r.get("kind") == "web":
            per_loc[loc]["web"] += 1
    for r in tests:
        if r.get("test_type") == "latency" and r.get("method") == "icmp":
            try:
                n = len(json.loads(r.get("raw_rtts_ms") or "[]"))
            except json.JSONDecodeError:
                n = 0
            per_loc[(_key(r)[0], _location(r))]["ping"] += n
    minimum = {
        "ftp_download": MIN_FTP_PER_LOCATION,
        "ftp_upload": MIN_FTP_PER_LOCATION,
        "web": MIN_WEB_PER_LOCATION,
        "ping": MIN_PINGS_PER_LOCATION,
    }
    regions: dict[str, set[str]] = defaultdict(set)
    for (region, loc), counts in sorted(per_loc.items()):
        regions[region].add(loc)
        for item, required in minimum.items():
            scope = f"{region} / {loc}"
            checks.append(SampleCheck(scope, f"{item} samples", required, counts.get(item, 0)))
    for region in sorted(regions):
        count = len(regions[region])
        checks.append(
            SampleCheck(region, "locations (Sch. 2 G13)", MIN_LOCATIONS_PER_REGION, count)
        )
    return checks


@dataclass
class Report:
    licensee: str
    period: str
    ftp_share_pct: float
    files: list[str]
    voice: list[VoiceRow]
    broadband: list[BroadbandRow]
    checks: list[SampleCheck]

    @property
    def passed(self) -> bool:
        results = [v.failure_ok for v in self.voice] + [v.drop_ok for v in self.voice]
        for b in self.broadband:
            results += [b.ftp_dl.ok, b.ftp_ul.ok, b.ping.ok, b.web.ok, b.rsrp.ok, b.rscp.ok]
        return all(r is not False for r in results) and all(c.ok for c in self.checks)


def build(exports: Exports, licensee: str, period: Period | None) -> Report:
    voice_rows = _in_period(exports.voice, period)
    qoe_rows = _in_period(exports.qoe, period)
    test_rows = _in_period(exports.tests, period)
    share = ftp_share_pct(period.start.year if period else None)
    voice = appendix_a(voice_rows)
    return Report(
        licensee=licensee,
        period=period.label if period else "all data",
        ftp_share_pct=share,
        files=exports.files,
        voice=voice,
        broadband=appendix_b(qoe_rows, test_rows, share),
        checks=sample_checks(voice, qoe_rows, test_rows),
    )


# ---------------------------------------------------------------- output


def _f(v: float | None, unit: str = "") -> str:
    return "–" if v is None else f"{v:.1f}{unit}"


def _ok(v: bool | None) -> str:
    return "n/a" if v is None else ("met" if v else "NOT MET")


def appendix_a_table(report: Report) -> list[list[object]]:
    head = [
        "region", "area_class", "call_attempts", "calls_failed", "call_failure_rate_pct",
        "failure_requirement", "calls_established", "calls_dropped", "call_drop_rate_pct",
        "drop_requirement", "mean_setup_time_s", "excluded_unknown", "technologies",
    ]  # fmt: skip
    rows: list[list[object]] = [head]
    for v in report.voice:
        rows.append([
            v.region, v.area_class, v.attempts, v.failed, v.failure_rate_pct,
            _ok(v.failure_ok), v.established, v.dropped, v.drop_rate_pct, _ok(v.drop_ok),
            v.mean_setup_s, v.unknown,
            "; ".join(f"{k} {n}" for k, n in sorted(v.technologies.items())),
        ])  # fmt: skip
    return rows


def appendix_b_table(report: Report) -> list[list[object]]:
    head = ["region", "area_class", "locations"]
    for name in ("ftp_dl", "ftp_ul", "ping", "web", "rsrp", "rscp"):
        head += [
            f"{name}_{col}" for col in ("samples", "meeting", "share_pct", "mean", "requirement")
        ]
    rows: list[list[object]] = [head]
    for b in report.broadband:
        row: list[object] = [b.region, b.area_class, len(b.locations)]
        for s in (b.ftp_dl, b.ftp_ul, b.ping, b.web, b.rsrp, b.rscp):
            row += [s.samples, s.meeting, s.share_pct, s.mean, _ok(s.ok)]
        rows.append(row)
    return rows


def checks_table(report: Report) -> list[list[object]]:
    rows: list[list[object]] = [["scope", "item", "required", "actual", "result"]]
    for c in report.checks:
        rows.append([c.scope, c.item, c.required, c.actual, "ok" if c.ok else "SHORT"])
    return rows


def write_csvs(report: Report, out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for name, table in (
        ("appendix_a_mobile_telephony.csv", appendix_a_table(report)),
        ("appendix_b_mobile_broadband.csv", appendix_b_table(report)),
        ("sample_checks.csv", checks_table(report)),
    ):
        path = out_dir / name
        with open(path, "w", newline="", encoding="utf-8") as f:
            csv.writer(f, lineterminator="\r\n").writerows(
                [["" if v is None else v for v in r] for r in table]
            )
        written.append(path)
    return written


def write_xlsx(report: Report, path: Path) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as e:  # pragma: no cover - depends on the environment
        raise ValueError("Excel output needs openpyxl: pip install 'onq-validation[xlsx]'") from e
    wb = Workbook()
    info = wb.active
    info.title = "Report"
    for row in (
        ["NICTA Telecommunications QoS Rule 2022 - measurement report (OpenNetIQ)"],
        ["Licensee", report.licensee],
        ["Measurement period", report.period],
        ["FTP share required", f"{report.ftp_share_pct:.0f} %"],
        ["Overall", "all requirements met" if report.passed else "requirements NOT met"],
        ["Generated", datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")],
        ["Source exports", len(report.files)],
    ):
        info.append(row)
    info["A1"].font = Font(bold=True)
    for title, table in (
        ("Appendix A", appendix_a_table(report)),
        ("Appendix B", appendix_b_table(report)),
        ("Sample checks", checks_table(report)),
    ):
        ws = wb.create_sheet(title)
        for r in table:
            ws.append(r)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "C2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def markdown(report: Report) -> str:
    lines = [
        f"# NICTA QoS Rule 2022 report: {report.licensee}",
        "",
        f"Period: {report.period} · source exports: {len(report.files)} · "
        f"FTP share required: {report.ftp_share_pct:.0f} %",
        f"**Overall: {'all requirements met' if report.passed else 'requirements NOT met'}**",
        "",
        "## Appendix A: mobile telephony (Schedule 1)",
    ]
    if not report.voice:
        lines.append("No call results.")
    else:
        lines += [
            "| Region | Area | Attempts | Failed | Call failure rate (≤ 2 %) | Established"
            " | Dropped | Call drop rate (≤ 2 %) | Mean setup |",
            "|---|---|---|---|---|---|---|---|---|",
        ]  # fmt: skip
        for v in report.voice:
            lines.append(
                f"| {v.region} | {v.area_class} | {v.attempts} | {v.failed} | "
                f"{_f(v.failure_rate_pct, ' %')} {_ok(v.failure_ok)} | {v.established} | "
                f"{v.dropped} | "
                f"{_f(v.drop_rate_pct, ' %')} {_ok(v.drop_ok)} | {_f(v.mean_setup_s, ' s')} |"
            )
    lines += ["", "## Appendix B: mobile broadband (Schedule 2)"]
    if not report.broadband:
        lines.append("No broadband results.")
    else:
        lines += [
            f"| Region | Area | FTP DL ≥ 2.5 Mbps (≥ {report.ftp_share_pct:.0f} %)"
            f" | FTP UL ≥ 0.5 Mbps (≥ {report.ftp_share_pct:.0f} %) | Ping ≤ 250 ms (≥ 80 %)"
            " | Web ≤ 10 s (≥ 80 %) | RSRP ≥ −110 dBm (≥ 80 %) | RSCP ≥ −100 dBm (≥ 80 %) |",
            "|---|---|---|---|---|---|---|---|",
        ]  # fmt: skip

        def cell(s: Share, unit: str) -> str:
            if s.samples == 0:
                return "no samples"
            return f"{_f(s.share_pct, ' %')} of {s.samples}, avg {_f(s.mean, unit)} {_ok(s.ok)}"

        for b in report.broadband:
            lines.append(
                f"| {b.region} | {b.area_class} | {cell(b.ftp_dl, ' Mbps')} | "
                f"{cell(b.ftp_ul, ' Mbps')} | {cell(b.ping, ' ms')} | {cell(b.web, ' s')} | "
                f"{cell(b.rsrp, ' dBm')} | {cell(b.rscp, ' dBm')} |"
            )
    short = [c for c in report.checks if not c.ok]
    lines += ["", "## Sample counts"]
    if not report.checks:
        lines.append("Nothing to check.")
    elif not short:
        lines.append(f"All {len(report.checks)} minimum sample counts reached.")
    else:
        lines += [f"{len(short)} of {len(report.checks)} minimum counts not reached:", ""]
        lines += [f"- {c.scope}: {c.item} {c.actual} / {c.required}" for c in short]
    return "\n".join(lines)
