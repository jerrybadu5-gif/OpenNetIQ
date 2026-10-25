import json
from datetime import timedelta

import pytest
from conftest import T0, iso, write_csv

from onq_validation import nicta
from onq_validation.cli import main

TAGS = ["region", "area_class", "location_name"]
VOICE = ["test_id", "timestamp", "kind", "outcome", "setup_ms", "voice_rat", *TAGS]
QOE = ["test_id", "timestamp", "kind", "status", "total_ms", "rate_mbps", "direction",
       "protocol", "serving_level_dbm", "serving_rat", *TAGS]  # fmt: skip
TESTS = ["test_type", "test_id", "timestamp", "method", "raw_rtts_ms", *TAGS]

SOUTH = ["Southern", "urban", "Waigani"]
HIGH = ["Highlands", "rural", "Banz"]


def t(i: int) -> str:
    return iso(T0 + timedelta(seconds=i))


@pytest.fixture
def exports(tmp_path):
    d = tmp_path / "exports"
    d.mkdir()
    calls = [[f"c{i}", t(i), "call", "completed", 5000, "VoLTE", *SOUTH] for i in range(97)]
    calls += [["d", t(97), "call", "dropped", 6000, "VoLTE>UMTS", *SOUTH]]
    calls += [[f"n{i}", t(98 + i), "call", "no_answer", None, "VoLTE", *SOUTH] for i in range(2)]
    calls += [["u", t(100), "call", "unknown", None, None, *SOUTH]]
    calls += [["s", t(101), "sms", "received", None, None, *SOUTH]]
    calls += [[f"h{i}", t(200 + i), "call", "completed", 7000, "GSM", *HIGH] for i in range(9)]
    calls += [["ht", t(209), "call", "setup_timeout", 25000, "GSM", *HIGH]]
    write_csv(d / "a_voice_sms.csv", VOICE, calls)

    qoe = [[f"f{i}", t(i), "file", "ok", 40000, 3.0, "download", "ftp", -100, "LTE", *SOUTH]
           for i in range(8)]  # fmt: skip
    qoe += [["f8", t(8), "file", "ok", 40000, 3.0, "download", "ftp", -112, "LTE", *SOUTH]]
    qoe += [["f9", t(9), "file", "failed", None, None, "download", "ftp", None, None, *SOUTH]]
    qoe += [[f"u{i}", t(20 + i), "file", "ok", 9000, 1.0, "upload", "ftp", -99, "LTE", *SOUTH]
            for i in range(10)]  # fmt: skip
    qoe += [["http", t(40), "file", "ok", 2000, 50.0, "download", "http", -90, "LTE", *SOUTH]]
    qoe += [[f"w{i}", t(50 + i), "web", "ok", 4000, None, None, None, None, None, *SOUTH]
            for i in range(9)]  # fmt: skip
    qoe += [["w9", t(59), "web", "ok", 12000, None, None, None, None, None, *SOUTH]]
    qoe += [["g", t(60), "file", "ok", 40000, 2.0, "download", "ftp", -98, "WCDMA", *HIGH]]
    write_csv(d / "a_qoe.csv", QOE, qoe)

    rtts = [100.0] * 85 + [300.0] * 10 + [None] * 5
    tests = [
        ["latency", "l1", t(70), "icmp", json.dumps(rtts), *SOUTH],
        ["latency", "l2", t(71), "tcp", json.dumps([1.0]), *SOUTH],
        ["speed", "s1", t(72), "http-mc-1.0", "", *SOUTH],
    ]
    write_csv(d / "a_tests.csv", TESTS, tests)
    (d / "a.csv").write_text("measurement_id\n")  # samples CSV: ignored
    return d


def test_appendix_a_and_b(exports):
    rep = nicta.build(nicta.find_exports([exports]), "Test Telco", None)
    south, high = rep.voice[1], rep.voice[0]  # sorted: Highlands, Southern
    assert (south.region, south.attempts, south.failed) == ("Southern", 100, 2)
    assert south.failure_rate_pct == 2.0 and south.failure_ok
    assert (south.established, south.dropped, south.drop_rate_pct) == (98, 1, 1.0)
    assert south.unknown == 1
    assert south.mean_setup_s == 5.0
    assert south.technologies == {"VoLTE": 99, "VoLTE>UMTS": 1}
    assert high.failure_rate_pct == 10.0 and high.failure_ok is False

    b = next(x for x in rep.broadband if x.region == "Southern")
    assert (b.ftp_dl.samples, b.ftp_dl.meeting, b.ftp_dl.share_pct) == (10, 9, 90.0)
    assert b.ftp_dl.ok  # 90 % required from 2024
    assert b.ftp_dl.mean == 3.0
    assert (b.ftp_ul.share_pct, b.ftp_ul.ok) == (100.0, True)
    assert (b.ping.samples, b.ping.share_pct, b.ping.ok) == (100, 85.0, True)
    assert b.ping.mean == 121.1  # (85 x 100 + 10 x 300) / 95; lost probes have no RTT
    assert (b.web.share_pct, b.web.ok, b.web.mean) == (90.0, True, 4.8)
    assert (b.rsrp.samples, b.rsrp.share_pct, b.rsrp.ok) == (9, 88.9, True)
    assert b.rscp.samples == 0 and b.rscp.ok is None
    hb = next(x for x in rep.broadband if x.region == "Highlands")
    assert hb.ftp_dl.ok is False  # 2.0 < 2.5 Mbps
    assert (hb.rscp.samples, hb.rscp.ok) == (1, True)

    short = {(c.scope, c.item) for c in rep.checks if not c.ok}
    assert ("Highlands", "call samples (Sch. 1 G11)") in short
    assert ("Southern", "locations (Sch. 2 G13)") in short
    assert ("Southern / Waigani", "ftp_download samples") not in short
    assert not rep.passed


def test_period_filter_and_share(exports):
    data = nicta.find_exports([exports])
    p = nicta.Period.parse("2026h2")
    assert p.label == "2026H2" and p.contains(T0)
    assert nicta.build(data, "x", p).voice
    empty = nicta.build(data, "x", nicta.Period.parse("2026H1"))
    assert empty.voice == [] and empty.broadband == [] and empty.passed
    assert nicta.ftp_share_pct(2023) == 80.0
    assert nicta.ftp_share_pct(2024) == 90.0
    assert nicta.build(data, "x", nicta.Period.parse("2023H1")).ftp_share_pct == 80.0
    with pytest.raises(ValueError):
        nicta.Period.parse("2026Q3")
    with pytest.raises(ValueError):
        nicta.find_exports([exports / "missing"])
    with pytest.raises(ValueError):
        nicta.find_exports([exports / "a.csv"])


def test_untagged_rows_and_helpers():
    rows = [{"kind": "call", "outcome": "completed", "setup_ms": ""}]
    v = nicta.appendix_a(rows)[0]
    assert (v.region, v.area_class) == (nicta.UNTAGGED, nicta.UNTAGGED)
    assert v.mean_setup_s is None
    assert nicta.pct(1, 0) is None and nicta.mean([]) is None
    bad = nicta.appendix_b([], [{"test_type": "latency", "method": "icmp", "raw_rtts_ms": "x"}], 90)
    assert bad[0].ping.samples == 0
    broken = [{"test_type": "latency", "method": "icmp", "raw_rtts_ms": "{"}]
    checks = nicta.sample_checks([], [], broken)
    assert any(c.item == "ping samples" and c.actual == 0 for c in checks)


def test_cli_markdown_csv_xlsx_json(exports, tmp_path, capsys):
    out = tmp_path / "out"
    code = main([
        "--json", str(tmp_path / "r.json"), "report", str(exports),
        "--licensee", "Test Telco", "--period", "2026H2",
        "--csv-dir", str(out), "--xlsx", str(out / "report.xlsx"),
    ])  # fmt: skip
    text = capsys.readouterr().out
    assert code == 1
    assert "# NICTA QoS Rule 2022 report: Test Telco" in text
    assert "2.0 % met" in text and "10.0 % NOT MET" in text
    assert "90.0 % of 10, avg 3.0 Mbps met" in text
    assert "Highlands: call samples (Sch. 1 G11) 10 / 100" in text
    a = (out / "appendix_a_mobile_telephony.csv").read_text().splitlines()
    assert a[0].startswith("region,area_class,call_attempts")
    assert any(line.startswith("Southern,urban,100,2,2.0,met") for line in a)
    assert (out / "appendix_b_mobile_broadband.csv").exists()
    assert "SHORT" in (out / "sample_checks.csv").read_text()
    from openpyxl import load_workbook

    wb = load_workbook(out / "report.xlsx")
    assert wb.sheetnames == ["Report", "Appendix A", "Appendix B", "Sample checks"]
    assert wb["Report"]["B2"].value == "Test Telco"
    assert json.loads((tmp_path / "r.json").read_text())["report"]["licensee"] == "Test Telco"
    assert main(["report", str(exports), "--period", "bad"]) == 2


def test_markdown_without_data():
    rep = nicta.Report("x", "all data", 90.0, [], [], [], [])
    text = nicta.markdown(rep)
    assert "No call results." in text and "No broadband results." in text
    assert "Nothing to check." in text
